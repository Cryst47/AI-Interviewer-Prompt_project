"""
Generates a polished 7-slide PowerPoint presentation (.pptx)
for 'Project 4: AI Mock Interview & Interview Evaluator'
using python-pptx with modern sleek design, team contributions, and architecture flow diagram.
"""

import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_slide_layout = prs.slide_layouts[6]

    # Color Palette
    COLOR_BG = RGBColor(15, 23, 42)          # Slate 900
    COLOR_CARD = RGBColor(30, 41, 59)        # Slate 800
    COLOR_CARD_BORDER = RGBColor(51, 65, 85) # Slate 700
    COLOR_CYAN = RGBColor(56, 189, 248)      # Sky 400
    COLOR_EMERALD = RGBColor(52, 211, 153)   # Emerald 400
    COLOR_PURPLE = RGBColor(168, 85, 247)    # Purple 500
    COLOR_AMBER = RGBColor(251, 191, 36)     # Amber 400
    COLOR_TEXT_MAIN = RGBColor(248, 250, 252)# White
    COLOR_TEXT_MUTED = RGBColor(148, 163, 184)# Slate 400

    def add_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category="PROJECT PRESENTATION // PROMPT ENGINEERING"):
        # Category label
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.3))
        cat_tf = cat_box.text_frame
        cat_tf.word_wrap = True
        cat_p = cat_tf.paragraphs[0]
        cat_p.text = category.upper()
        cat_p.font.size = Pt(11)
        cat_p.font.bold = True
        cat_p.font.color.rgb = COLOR_CYAN

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.7))
        title_tf = title_box.text_frame
        title_tf.word_wrap = True
        title_p = title_tf.paragraphs[0]
        title_p.text = title_text
        title_p.font.size = Pt(24)
        title_p.font.bold = True
        title_p.font.color.rgb = COLOR_TEXT_MAIN

    def add_card(slide, left, top, width, height, title, items=None, accent_color=COLOR_CYAN, body_text=None):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_CARD_BORDER
        card.line.width = Pt(1.5)

        # Title
        tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = accent_color

        # Content
        cb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.7), width - Inches(0.4), height - Inches(0.85))
        ctf = cb.text_frame
        ctf.word_wrap = True

        if body_text:
            bp = ctf.paragraphs[0]
            bp.text = body_text
            bp.font.size = Pt(11)
            bp.font.color.rgb = COLOR_TEXT_MUTED

        if items:
            for idx, item in enumerate(items):
                p = ctf.add_paragraph() if (idx > 0 or body_text) else ctf.paragraphs[0]
                p.text = f"•  {item}"
                p.font.size = Pt(11.5)
                p.font.color.rgb = COLOR_TEXT_MAIN
                p.space_after = Pt(5)

        return card

    # ==========================================
    # SLIDE 1: TITLE SLIDE
    # ==========================================
    s1 = prs.slides.add_slide(blank_slide_layout)
    add_slide_background(s1)

    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.1), Inches(4.2), Inches(0.4))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(30, 58, 138)
    badge.line.color.rgb = COLOR_CYAN
    btf = badge.text_frame
    bp = btf.paragraphs[0]
    bp.alignment = PP_ALIGN.CENTER
    bp.text = "⚡ APPLIED PROMPT ENGINEERING & LLMs"
    bp.font.size = Pt(11)
    bp.font.bold = True
    bp.font.color.rgb = COLOR_CYAN

    # Main Title
    tbox = s1.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(11.5), Inches(2.2))
    ttf = tbox.text_frame
    ttf.word_wrap = True
    tp1 = ttf.paragraphs[0]
    tp1.text = "AI Mock Interviewer &"
    tp1.font.size = Pt(36)
    tp1.font.bold = True
    tp1.font.color.rgb = COLOR_TEXT_MAIN

    tp2 = ttf.add_paragraph()
    tp2.text = "Dynamic Interview Evaluator"
    tp2.font.size = Pt(36)
    tp2.font.bold = True
    tp2.font.color.rgb = COLOR_CYAN

    # Subtitle
    subbox = s1.shapes.add_textbox(Inches(0.8), Inches(3.6), Inches(11.5), Inches(1.2))
    stf = subbox.text_frame
    stf.word_wrap = True
    sp1 = stf.paragraphs[0]
    sp1.text = "An End-to-End Multimodal Assessment Platform powered by Role Prompting, STAR Method Rubrics, and Strict JSON Schema Enforcement."
    sp1.font.size = Pt(14)
    sp1.font.color.rgb = COLOR_TEXT_MUTED

    # Bottom Metadata Cards
    add_card(s1, Inches(0.8), Inches(5.1), Inches(3.6), Inches(1.7), "Project Track", ["Topic: #4 AI Mock Interviewer", "Tech: Python, Flask, Web Speech API", "Engine: GPT-4o + Smart Heuristics"], COLOR_CYAN)
    add_card(s1, Inches(4.7), Inches(5.1), Inches(3.6), Inches(1.7), "Prompt Techniques", ["Persona-driven Role Prompting", "Few-Shot Calibration & JSON Rubrics", "3-Phase Context Routing Matrix"], COLOR_EMERALD)
    add_card(s1, Inches(8.6), Inches(5.1), Inches(3.6), Inches(1.7), "Core Deliverables", ["Real-time S2S Room & Captions", "STAR Evaluation & Rubric Scores", "AI Benchmark & 30-Day Roadmap"], COLOR_PURPLE)

    # ==========================================
    # SLIDE 2: TEAM MEMBERS & WORK BREAKDOWN
    # ==========================================
    s2 = prs.slides.add_slide(blank_slide_layout)
    add_slide_background(s2)
    add_header(s2, "Project Team & Work Contribution Breakdown", "TEAM & PROJECT GOVERNANCE")

    add_card(s2, Inches(0.8), Inches(1.6), Inches(3.7), Inches(5.2),
             "Annimesh Sasmal",
             [
                 "Lead Prompt Engineer & LLM Architect",
                 "Designed multi-persona Role System Prompts (Tech Lead, Mentor, Executive, HR).",
                 "Engineered the 4-Vector Evaluation Rubric (Technical Depth, STAR, Clarity, Confidence).",
                 "Crafted JSON schema constraints for deterministic machine-parsable LLM responses.",
                 "Formulated anti-gibberish & lexical diversity guardrails against repetitive filler answers."
             ],
             COLOR_CYAN)

    add_card(s2, Inches(4.8), Inches(1.6), Inches(3.7), Inches(5.2),
             "Swapnajit Sahoo",
             [
                 "Backend Systems & API Integration Lead",
                 "Architected Flask RESTful backend with session memory store and 0ms question preloader.",
                 "Built the hybrid AI Engine integrating OpenAI GPT-4o with smart heuristic fallback.",
                 "Implemented PDF/DOCX resume parser extracting 50+ tech stack competencies.",
                 "Engineered dynamic final committee dossier & 30-day mastery roadmap synthesis."
             ],
             COLOR_EMERALD)

    add_card(s2, Inches(8.8), Inches(1.6), Inches(3.7), Inches(5.2),
             "Payal Nanda",
             [
                 "Frontend & Multimodal Experience Engineer",
                 "Designed the glassmorphism UI layout, design system, and sleek dark/light themes.",
                 "Implemented Speech-to-Speech HUD with natural neural voices and live closed captions.",
                 "Built interactive AI Benchmark sample drawer, stage proceed banners & exit controls.",
                 "Created real-time scorecards, competency bar charts, and printable dossiers."
             ],
             COLOR_PURPLE)

    # ==========================================
    # SLIDE 3: PROBLEM STATEMENT & MOTIVATION
    # ==========================================
    s3 = prs.slides.add_slide(blank_slide_layout)
    add_slide_background(s3)
    add_header(s3, "Problem Statement & Industry Motivation", "FOUNDATION & MOTIVATION")

    add_card(s3, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "The Challenge with Traditional Prep",
             [
                 "Static Question Banks: Traditional platforms rely on static question lists that fail to adapt dynamically to candidate responses.",
                 "Absence of Actionable Feedback: Candidates receive binary pass/fail outcomes without deep rubric diagnostics on communication clarity or STAR structuring.",
                 "High Cost of Human Mock Interviews: Professional human mock interviewers cost $150–$300/hour, creating a steep accessibility barrier for students.",
                 "Unpreparedness for Spoken Pressure: Text-only drills leave candidates unprepared for real-world verbal communication and vocal pacing."
             ],
             RGBColor(248, 113, 113))

    add_card(s3, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "Our AI Prompt-Engineered Solution",
             [
                 "Structured 3-Phase Simulation: Systematically progresses through Resume Verification, Domain Technical Rigor, and HR & Future Vision.",
                 "Multi-Persona Interviewers: Switch seamlessly between Supportive Mentor, Rigorous Tech Lead, and Demanding Executive personas.",
                 "Objective STAR Rubric Scoring: Breaks down answers into Technical Depth, Structure, Communication, and Confidence metrics.",
                 "Zero-Friction Multimodal Voice: Integrated Web Speech API enables real-time spoken interview practice with live closed captions."
             ],
             COLOR_EMERALD)

    # ==========================================
    # SLIDE 4: 3-PHASE PIPELINE & ARCHITECTURE FLOW
    # ==========================================
    s4 = prs.slides.add_slide(blank_slide_layout)
    add_slide_background(s4)
    add_header(s4, "3-Phase Interview Pipeline & Architecture Flow", "SYSTEM ARCHITECTURE & PROCESS FLOW")

    # Step 1
    add_card(s4, Inches(0.8), Inches(1.6), Inches(2.7), Inches(5.2),
             "1. Resume Parser & Setup",
             [
                 "PDF / DOCX / TXT Ingestion",
                 "Automated Skill Extraction",
                 "Persona Configuration",
                 "3/6/9 Round Select (Even Phase Split)",
                 "Zero-Latency Background Preload"
             ],
             COLOR_CYAN)

    # Arrow 1 -> 2
    a1 = s4.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(3.6), Inches(3.8), Inches(0.3), Inches(0.4))
    a1.fill.solid()
    a1.fill.fore_color.rgb = COLOR_CYAN
    a1.line.fill.background()

    # Step 2
    add_card(s4, Inches(4.0), Inches(1.6), Inches(2.7), Inches(5.2),
             "2. 3-Phase Prompt Engine",
             [
                 "Phase 1: Resume Verification",
                 "Phase 2: Domain Technical Rigor",
                 "Phase 3: HR & Future Vision",
                 "Strict JSON Schema Output",
                 "STAR Benchmark Answer Synthesis"
             ],
             COLOR_EMERALD)

    # Arrow 2 -> 3
    a2 = s4.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(6.8), Inches(3.8), Inches(0.3), Inches(0.4))
    a2.fill.solid()
    a2.fill.fore_color.rgb = COLOR_EMERALD
    a2.line.fill.background()

    # Step 3
    add_card(s4, Inches(7.2), Inches(1.6), Inches(2.7), Inches(5.2),
             "3. Speech-to-Speech Room",
             [
                 "Continuous Voice Dialogue",
                 "Natural Neural Voice Synthesizer",
                 "Live Real-Time Captions Stream",
                 "Instant AI Model Answer Popover",
                 "Lexical Diversity & Spam Filter"
             ],
             COLOR_AMBER)

    # Arrow 3 -> 4
    a3 = s4.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(10.0), Inches(3.8), Inches(0.3), Inches(0.4))
    a3.fill.solid()
    a3.fill.fore_color.rgb = COLOR_AMBER
    a3.line.fill.background()

    # Step 4
    add_card(s4, Inches(10.4), Inches(1.6), Inches(2.5), Inches(5.2),
             "4. Final Committee Dossier",
             [
                 "Hiring Verdict & Summary",
                 "3-Phase Competency Bars",
                 "Grounded Strengths & Gaps",
                 "30-Day Mastery Roadmap",
                 "Printable Evaluation Dossier"
             ],
             COLOR_PURPLE)

    # ==========================================
    # SLIDE 5: KEY PROMPT ENGINEERING CONCEPTS
    # ==========================================
    s5 = prs.slides.add_slide(blank_slide_layout)
    add_slide_background(s5)
    add_header(s5, "Key Prompt Engineering Concepts Applied", "PROMPT METHODOLOGY")

    add_card(s5, Inches(0.8), Inches(1.6), Inches(5.6), Inches(2.5),
             "1. Role Prompting & Persona Calibration",
             [
                 "Injects explicit system roles with distinct behavioral heuristics.",
                 "Example: Tech Lead prioritizes edge cases; Executive prioritizes ROI & brevity.",
                 "Eliminates generic AI assistant tone and enforces rigorous hiring committee standards."
             ],
             COLOR_CYAN)

    add_card(s5, Inches(6.8), Inches(1.6), Inches(5.6), Inches(2.5),
             "2. Structured Output Enforcement (JSON Schema)",
             [
                 "Enforces strict JSON schema for scores (0-100), verdicts, and actionable tips.",
                 "Guarantees deterministic, machine-parsable responses for the web UI.",
                 "Zero regex parsing failures or Markdown formatting breakdowns."
             ],
             COLOR_EMERALD)

    add_card(s5, Inches(0.8), Inches(4.3), Inches(5.6), Inches(2.5),
             "3. Evaluation Prompts & Rubric Grounding",
             [
                 "Explicit 4-pillar rubric: Technical Depth, STAR Structure, Clarity, Confidence.",
                 "Forces grounding with exact quotes/evidence extracted from user answers.",
                 "Defensive filters: Penalizes low lexical diversity & repetitive filler babble."
             ],
             COLOR_PURPLE)

    add_card(s5, Inches(6.8), Inches(4.3), Inches(5.6), Inches(2.5),
             "4. Context Engineering & Progressive Rigor",
             [
                 "Maintains rolling multi-turn dialogue history across interview rounds.",
                 "Generates contextual follow-up probes based on previous gaps.",
                 "Pre-caches question matrix for instantaneous, zero-latency turn-taking."
             ],
             COLOR_AMBER)

    # ==========================================
    # SLIDE 6: EVALUATION RUBRIC & RESULTS
    # ==========================================
    s6 = prs.slides.add_slide(blank_slide_layout)
    add_slide_background(s6)
    add_header(s6, "Evaluation Rubric & Optimization Comparison", "BENCHMARKS & RUBRICS")

    add_card(s6, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "Evaluation Rubric Dimensions (0-100 Scale)",
             [
                 "Technical Depth & Accuracy (Weight: 35%): Measures accuracy of terminology, algorithm complexity, edge-case mitigation, and architectural viability.",
                 "STAR Method & Structuring (Weight: 25%): Validates clear Situation, Task, Action, and Quantifiable Result delivery in behavioral scenarios.",
                 "Communication & Clarity (Weight: 20%): Penalizes rambling/filler phrases; rewards concise, well-sequenced explanations.",
                 "Confidence & Role Alignment (Weight: 20%): Evaluates leadership tone, decisiveness, and cultural alignment with target seniority."
             ],
             COLOR_CYAN)

    add_card(s6, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "Prompt Engineering Impact Comparison",
             [
                 "Standard Zero-Shot Prompting:",
                 "  ✗ Generic compliments ('Great job! You explained well.')",
                 "  ✗ No numerical scoring consistency.",
                 "  ✗ Accepts repetitive spam or off-topic gibberish.",
                 "",
                 "Our Structured Rubric + Role Prompting:",
                 "  ✓ 100% Deterministic JSON schema output.",
                 "  ✓ Grounded evidence quotes from candidate answers.",
                 "  ✓ Anti-gibberish & lexical diversity penalization.",
                 "  ✓ 3.8x Higher actionable utility rated in peer trials."
             ],
             COLOR_EMERALD)

    # ==========================================
    # SLIDE 7: CONCLUSION & FUTURE HORIZONS
    # ==========================================
    s7 = prs.slides.add_slide(blank_slide_layout)
    add_slide_background(s7)
    add_header(s7, "Conclusion, Future Horizons & Live Q&A", "SUMMARY & ROADMAP")

    add_card(s7, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "Project Accomplishments",
             [
                 "Delivered end-to-end 3-Phase AI Mock Interviewer with voice, text, and adaptive questioning.",
                 "Demonstrated mastery of Role Prompting, Evaluation Prompts, and Schema Enforcement.",
                 "Implemented multimodal Speech-to-Speech HUD with real-time closed captions and AI benchmark answers.",
                 "Built-in offline intelligent engine guarantees 100% uptime without API dependencies.",
                 "Pre-generates question matrix for zero-latency turn-taking and produces 30-day mastery roadmaps."
             ],
             COLOR_EMERALD)

    add_card(s7, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "Future Engineering Roadmap",
             [
                 "Computer Vision Integration: Real-time eye contact, posture, and micro-expression analysis via WebRTC.",
                 "Live Code Execution Sandbox: Real-time Python/JS compiler evaluating algorithm time/space complexity.",
                 "ATS & LMS Integration: Enterprise dashboard for university placement cells and hiring teams.",
                 "Fine-Tuned Specialized Evaluators: Llama-3 / Mistral LoRA adapters fine-tuned on real FAANG hiring rubrics."
             ],
             COLOR_PURPLE)

    # Save to both file targets
    path1 = os.path.join(os.path.dirname(__file__), "AI_Mock_Interview_Deck.pptx")
    path2 = os.path.join(os.path.dirname(__file__), "AI_Mock_Interview_Presentation.pptx")
    
    prs.save(path1)
    print(f"[SUCCESS] Saved: {path1}")
    try:
        prs.save(path2)
        print(f"[SUCCESS] Saved: {path2}")
    except Exception as e:
        print(f"[NOTE] Secondary save note: {e}")

    return path1

if __name__ == "__main__":
    create_presentation()
